"""Coordinate normalization endpoints (STEP 1)."""

from __future__ import annotations

import csv
import io
import re
import zipfile

from fastapi import APIRouter, File, HTTPException, UploadFile

from app.schemas.models import NormalizeRequest, NormalizeResponse
from app.services import boundary
from app.services import normalizer as nz
from app.services import url_resolver

router = APIRouter(prefix="/api/coordinates", tags=["coordinates"])


def _apply_territory(result: nz.NormalizationResult) -> dict:
    """Drop coordinates outside Uzbekistan, moving them to ``invalid`` and
    reporting how many were removed via ``outside_territory``."""
    inside: list[nz.ParsedCoordinate] = []
    outside: list[nz.ParsedCoordinate] = []
    for c in result.coordinates:
        if boundary.is_in_uzbekistan(c.latitude, c.longitude):
            inside.append(c)
        else:
            c.valid = False
            c.error = boundary.OUTSIDE_MESSAGE
            outside.append(c)

    for i, c in enumerate(inside, start=1):
        c.point_number = i

    result.coordinates = inside
    result.invalid.extend(outside)

    data = result.to_dict()
    data["outside_territory"] = len(outside)
    return data


@router.post("/normalize", response_model=NormalizeResponse)
async def normalize(req: NormalizeRequest):
    """Normalize free-form text, CSV rows, or map URLs into standard coords."""
    text = req.text or ""
    extra_pairs: list[tuple[float, float]] = []

    if req.resolve_urls and url_resolver.is_url(text):
        try:
            extra_pairs = await url_resolver.resolve_text(text)
        except Exception:
            # URL resolution is best-effort and must never fail the request.
            extra_pairs = []

    try:
        result = nz.normalize_text(text)
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=422,
                            detail=f"Could not parse coordinates: {exc}")

    if extra_pairs:
        # Merge URL-resolved coordinates, re-running duplicate detection.
        merged = [(c.latitude, c.longitude) for c in result.coordinates] + extra_pairs
        result = nz.normalize_pairs(merged)

    return _apply_territory(result)


@router.post("/upload", response_model=NormalizeResponse)
async def upload_file(file: UploadFile = File(...)):
    """Normalize coordinates from an uploaded TXT / CSV / XLSX / KML / KMZ file."""
    raw = await file.read()
    name = (file.filename or "").lower()

    if name.endswith(".xlsx") or name.endswith(".xls"):
        text = _xlsx_to_text(raw)
    elif name.endswith(".kmz"):
        text = _kmz_to_text(raw)
    elif name.endswith(".kml"):
        text = _kml_to_text(raw)
    else:
        text = raw.decode("utf-8", errors="replace")
        # Normalize CSV delimiters to a parser-friendly form.
        text = _csv_to_text(text)

    result = nz.normalize_text(text)
    return _apply_territory(result)


def _csv_to_text(text: str) -> str:
    try:
        dialect = csv.Sniffer().sniff(text[:2048], delimiters=",;\t")
    except csv.Error:
        return text
    out_lines = []
    reader = csv.reader(io.StringIO(text), dialect)
    for row in reader:
        out_lines.append(" ".join(cell.strip() for cell in row if cell.strip()))
    return "\n".join(out_lines)


_COORDINATES_RE = re.compile(
    r"<coordinates[^>]*>(.*?)</coordinates>", re.DOTALL | re.IGNORECASE)


def _kmz_to_text(raw: bytes) -> str:
    """Extract the embedded KML from a KMZ archive and parse its coordinates."""
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as zf:
            kml_name = next(
                (n for n in zf.namelist() if n.lower().endswith(".kml")), None)
            if kml_name is None:
                raise HTTPException(
                    status_code=422,
                    detail="KMZ ichida .kml fayl topilmadi")
            kml_bytes = zf.read(kml_name)
    except zipfile.BadZipFile:
        raise HTTPException(status_code=422, detail="Yaroqsiz KMZ fayl")
    return _kml_to_text(kml_bytes)


def _kml_to_text(raw: bytes | str) -> str:
    """Convert KML ``<coordinates>`` blocks into ``lat lon`` lines.

    KML stores coordinate tuples as ``lon,lat[,alt]`` separated by whitespace.
    We emit ``lat lon`` per line so the normalizer (which assumes lat/lon order
    when no hemisphere marker is present) reads them correctly.
    """
    text = raw.decode("utf-8", errors="replace") if isinstance(raw, bytes) else raw
    lines: list[str] = []
    for block in _COORDINATES_RE.findall(text):
        for tuple_str in block.split():
            parts = tuple_str.strip().split(",")
            if len(parts) >= 2 and parts[0] and parts[1]:
                lon, lat = parts[0], parts[1]
                lines.append(f"{lat} {lon}")
    return "\n".join(lines)


def _xlsx_to_text(raw: bytes) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    lines = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                lines.append(" ".join(cells))
    return "\n".join(lines)
