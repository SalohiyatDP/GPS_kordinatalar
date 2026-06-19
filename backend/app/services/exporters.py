"""Export module (STEP 2 & STEP 5).

Generates XLSX, KML, KMZ, GeoJSON and PDF outputs. All functions return raw
``bytes`` so they can be streamed directly from the API.
"""

from __future__ import annotations

import io
import json
import zipfile
from datetime import datetime

import simplekml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (Image, Paragraph, SimpleDocTemplate, Spacer,
                                Table, TableStyle)

from .normalizer import decimal_to_dms

# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(bold=True, color="FFFFFF")


def _style_header(ws, ncols: int):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")


def points_to_xlsx(points: list[dict]) -> bytes:
    """points.xlsx — Point Number, Latitude, Longitude, DMS."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Points"
    headers = ["Point Number", "Latitude", "Longitude", "DMS"]
    ws.append(headers)
    _style_header(ws, len(headers))
    for p in points:
        ws.append([
            p["point_number"],
            p["latitude"],
            p["longitude"],
            p.get("dms", ""),
        ])
    widths = [14, 16, 16, 34]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    return _wb_bytes(wb)


def analysis_to_xlsx(points: list[dict], area: dict, perimeter: dict,
                     contours: list[dict] | None = None,
                     summary: str = "") -> bytes:
    """Analysis.xlsx — full cadastral report workbook (multiple sheets)."""
    wb = Workbook()

    ws = wb.active
    ws.title = "Coordinates"
    headers = ["Point Number", "Latitude", "Longitude", "DMS"]
    ws.append(headers)
    _style_header(ws, len(headers))
    for p in points:
        ws.append([p["point_number"], p["latitude"], p["longitude"], p.get("dms", "")])
    for i, w in enumerate([14, 16, 16, 34], start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    ws2 = wb.create_sheet("Summary")
    ws2.append(["Metric", "Value"])
    _style_header(ws2, 2)
    ws2.append(["Area (m²)", area.get("square_meters")])
    ws2.append(["Area (hectares)", area.get("hectares")])
    ws2.append(["Area (km²)", area.get("square_kilometers")])
    ws2.append(["Area (sotix)", area.get("sotix")])
    ws2.append(["Perimeter (m)", perimeter.get("meters")])
    ws2.append(["Perimeter (km)", perimeter.get("kilometers")])
    ws2.append(["Generated", datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC")])
    if summary:
        ws2.append(["Summary", summary])
    ws2.column_dimensions["A"].width = 22
    ws2.column_dimensions["B"].width = 50

    if contours:
        ws3 = wb.create_sheet("Contours")
        c_headers = ["Contour", "Region", "District", "Massif",
                     "Contour Area (m²)", "Intersection Area (m²)",
                     "Coverage %", "Status"]
        ws3.append(c_headers)
        _style_header(ws3, len(c_headers))
        for c in contours:
            ws3.append([
                c.get("contour"),
                c.get("region"),
                c.get("district"),
                c.get("massif"),
                c.get("contour_area"),
                c.get("intersection_area"),
                c.get("coverage_percent"),
                c.get("status"),
            ])
        for i, w in enumerate([12, 16, 16, 16, 18, 20, 12, 10], start=1):
            ws3.column_dimensions[ws3.cell(row=1, column=i).column_letter].width = w

    return _wb_bytes(wb)


def _wb_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# KML / KMZ
# ---------------------------------------------------------------------------

def _kml_to_kmz_bytes(kml: simplekml.Kml) -> bytes:
    """Package a simplekml document into in-memory KMZ (zip) bytes."""
    kml_str = kml.kml()
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("doc.kml", kml_str)
    return buf.getvalue()

def points_to_kmz(points: list[dict]) -> bytes:
    """KMZ with one placemark per point."""
    kml = simplekml.Kml()
    for p in points:
        pnt = kml.newpoint(name=str(p["point_number"]))
        pnt.coords = [(p["longitude"], p["latitude"])]
        pnt.description = (
            f"Point {p['point_number']}\n"
            f"{p.get('dms', '')}\n"
            f"{p['latitude']}, {p['longitude']}"
        )
    return _kml_to_kmz_bytes(kml)


def polygon_to_kmz(points: list[dict], name: str = "Polygon") -> bytes:
    """polygon.kmz — closed polygon compatible with Google Earth."""
    kml = simplekml.Kml()
    ring = [(p["longitude"], p["latitude"]) for p in points]
    if ring and ring[0] != ring[-1]:
        ring.append(ring[0])
    pol = kml.newpolygon(name=name, outerboundaryis=ring)
    pol.style.linestyle.color = simplekml.Color.red
    pol.style.linestyle.width = 3
    pol.style.polystyle.color = simplekml.Color.changealphaint(80, simplekml.Color.red)
    # Add points too for reference.
    for p in points:
        pnt = kml.newpoint(name=str(p["point_number"]))
        pnt.coords = [(p["longitude"], p["latitude"])]
    return _kml_to_kmz_bytes(kml)


def analysis_to_kmz(points: list[dict], contours: list[dict] | None = None) -> bytes:
    """Analysis.kmz — polygon + points + contour intersections colored."""
    kml = simplekml.Kml()
    # User polygon (red).
    ring = [(p["longitude"], p["latitude"]) for p in points]
    if ring and ring[0] != ring[-1]:
        ring.append(ring[0])
    if len(ring) >= 4:
        pol = kml.newpolygon(name="User Polygon", outerboundaryis=ring)
        pol.style.linestyle.color = simplekml.Color.red
        pol.style.linestyle.width = 3
        pol.style.polystyle.color = simplekml.Color.changealphaint(60, simplekml.Color.red)

    pts_folder = kml.newfolder(name="Points")
    for p in points:
        pnt = pts_folder.newpoint(name=str(p["point_number"]))
        pnt.coords = [(p["longitude"], p["latitude"])]

    if contours:
        c_folder = kml.newfolder(name="Contours")
        for c in contours:
            geojson = c.get("geometry")
            if not geojson:
                continue
            full = c.get("status") == "Full"
            color = simplekml.Color.green if full else simplekml.Color.yellow
            for ring_coords in _iter_polygon_rings(geojson):
                cpol = c_folder.newpolygon(
                    name=str(c.get("contour")),
                    outerboundaryis=ring_coords,
                )
                cpol.style.polystyle.color = simplekml.Color.changealphaint(90, color)
                cpol.style.linestyle.color = color
    return _kml_to_kmz_bytes(kml)


def _iter_polygon_rings(geojson: dict):
    gtype = geojson.get("type")
    coords = geojson.get("coordinates", [])
    if gtype == "Polygon":
        if coords:
            yield [(x, y) for x, y in coords[0]]
    elif gtype == "MultiPolygon":
        for poly in coords:
            if poly:
                yield [(x, y) for x, y in poly[0]]


# ---------------------------------------------------------------------------
# GeoJSON
# ---------------------------------------------------------------------------

def to_geojson(points: list[dict], polygon_geojson: dict | None = None,
               contours: list[dict] | None = None) -> bytes:
    features = []
    for p in points:
        features.append({
            "type": "Feature",
            "geometry": {"type": "Point", "coordinates": [p["longitude"], p["latitude"]]},
            "properties": {"point_number": p["point_number"], "dms": p.get("dms", "")},
        })
    if polygon_geojson:
        features.append({
            "type": "Feature",
            "geometry": polygon_geojson,
            "properties": {"name": "User Polygon"},
        })
    if contours:
        for c in contours:
            if c.get("geometry"):
                features.append({
                    "type": "Feature",
                    "geometry": c["geometry"],
                    "properties": {k: v for k, v in c.items() if k != "geometry"},
                })
    fc = {"type": "FeatureCollection", "features": features}
    return json.dumps(fc, ensure_ascii=False, indent=2).encode("utf-8")


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def to_pdf(points: list[dict], area: dict, perimeter: dict,
           contours: list[dict] | None = None,
           summary: str = "", title: str = "Cadastre Analysis Report",
           map_image_png: bytes | None = None) -> bytes:
    """Analysis.pdf — coordinate table, area, perimeter, optional map & contours."""
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=1.5 * cm,
                            bottomMargin=1.5 * cm, leftMargin=1.5 * cm,
                            rightMargin=1.5 * cm)
    styles = getSampleStyleSheet()
    elements = []

    elements.append(Paragraph(title, styles["Title"]))
    elements.append(Paragraph(
        f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}",
        styles["Normal"]))
    elements.append(Spacer(1, 0.4 * cm))

    if map_image_png:
        try:
            img = Image(io.BytesIO(map_image_png))
            img._restrictSize(17 * cm, 9 * cm)
            elements.append(img)
            elements.append(Spacer(1, 0.4 * cm))
        except Exception:
            pass

    # Coordinate table.
    elements.append(Paragraph("Coordinates", styles["Heading2"]))
    data = [["#", "Latitude", "Longitude", "DMS"]]
    for p in points:
        data.append([
            str(p["point_number"]),
            f"{p['latitude']:.6f}",
            f"{p['longitude']:.6f}",
            p.get("dms", ""),
        ])
    table = Table(data, hAlign="LEFT", colWidths=[1.2 * cm, 3.2 * cm, 3.2 * cm, 7 * cm])
    table.setStyle(_table_style())
    elements.append(table)
    elements.append(Spacer(1, 0.4 * cm))

    # Area & perimeter.
    elements.append(Paragraph("Geometry", styles["Heading2"]))
    geo = [
        ["Area (m²)", f"{area.get('square_meters', 0):,.2f}"],
        ["Area (hectares)", f"{area.get('hectares', 0):,.4f}"],
        ["Area (km²)", f"{area.get('square_kilometers', 0):,.6f}"],
        ["Area (sotix)", f"{area.get('sotix', 0):,.2f}"],
        ["Perimeter (m)", f"{perimeter.get('meters', 0):,.2f}"],
        ["Perimeter (km)", f"{perimeter.get('kilometers', 0):,.4f}"],
    ]
    gtable = Table(geo, hAlign="LEFT", colWidths=[6 * cm, 8 * cm])
    gtable.setStyle(_table_style())
    elements.append(gtable)
    elements.append(Spacer(1, 0.4 * cm))

    if contours:
        elements.append(Paragraph("Contour Analysis", styles["Heading2"]))
        cdata = [["Contour", "Region", "District", "Massif",
                  "Area m²", "Intersect m²", "Cov %", "Status"]]
        for c in contours:
            cdata.append([
                str(c.get("contour", "")),
                str(c.get("region", "")),
                str(c.get("district", "")),
                str(c.get("massif", "")),
                f"{c.get('contour_area', 0):,.0f}",
                f"{c.get('intersection_area', 0):,.0f}",
                f"{c.get('coverage_percent', 0):.1f}",
                str(c.get("status", "")),
            ])
        ctable = Table(cdata, hAlign="LEFT")
        ctable.setStyle(_table_style())
        elements.append(ctable)
        elements.append(Spacer(1, 0.4 * cm))

    if summary:
        elements.append(Paragraph("Summary", styles["Heading2"]))
        elements.append(Paragraph(summary, styles["Normal"]))

    doc.build(elements)
    return buf.getvalue()


def _table_style() -> TableStyle:
    return TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EEF3F8")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ])
